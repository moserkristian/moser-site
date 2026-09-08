#!/usr/bin/env python3
"""Build a single-column .docx CV. Paragraphs only — no tables or text boxes."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, Mm, RGBColor


def set_run_font(run, size_pt, bold=False):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Calibri")
    rfonts.set(qn("w:hAnsi"), "Calibri")
    rfonts.set(qn("w:eastAsia"), "Calibri")
    rfonts.set(qn("w:cs"), "Calibri")


def add_para(doc, text, size=11, bold=False, space_after=4, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size, bold)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.clear()
    run = p.add_run(text)
    set_run_font(run, 11, False)


def main():
    out = Path(__file__).resolve().parent / "Kristian-Moser-CV.docx"
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    section.top_margin = Mm(14)
    section.bottom_margin = Mm(14)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    add_para(doc, "Kristián Moser", 18, True, space_after=2)
    add_para(doc, "Senior .NET engineer", 12, True, space_after=4)
    add_para(doc, "Bratislava · moser.kristian@gmail.com · +421 918 829 553", 10)
    add_para(doc, "linkedin.com/in/moserkristian · github.com/moserkristian", 10, space_after=8)

    add_para(doc, "Profile", 12, True, space_before=6, space_after=4)
    add_para(
        doc,
        "Senior .NET engineer (C#, ASP.NET Core, Blazor). Product and client work in laboratory APIs, aviation data, and industrial cloud. Open to contract work through Moser Consulting s.r.o.",
        11,
        space_after=8,
    )

    add_para(doc, "Technical skills", 12, True, space_before=6, space_after=4)
    add_para(doc, "Languages: C#, SQL, JavaScript", 11)
    add_para(doc, "Backend: ASP.NET Core, REST APIs, Entity Framework Core, Blazor", 11)
    add_para(doc, "Delivery: Azure DevOps, Docker, CI/CD", 11)
    add_para(doc, "Cloud: Microsoft Azure", 11)
    add_para(doc, "Data: SQL Server, PostgreSQL, PostGIS, T-SQL / SQL", 11, space_after=8)

    add_para(doc, "Experience", 12, True, space_before=6, space_after=4)

    add_para(doc, "Senior .NET engineer — Moser Consulting s.r.o.", 11, True, space_before=8)
    add_para(doc, "Aug 2024 – Present · Bratislava / Remote", 10)

    add_para(doc, "NG Aviation", 11, True, space_before=8)
    add_para(doc, "Mar 2025 – Oct 2025 · Bratislava · Hybrid", 10)
    add_bullet(doc, "NOTAM ingestion to AIXM JSON and publish to a cloud database.")
    add_bullet(doc, "REST APIs over aerospace data.")
    add_bullet(doc, "PostgreSQL / PostGIS for geospatial queries.")
    add_bullet(doc, "Docker and CI/CD.")
    add_bullet(doc, "Cleaned up layering in existing services.")

    add_para(doc, "ComAp Group", 11, True, space_before=8)
    add_para(doc, "Aug 2024 – Feb 2025 · Prague · Hybrid", 10)
    add_bullet(doc, "Early cloud platform work: service boundaries, inter-service APIs, Clean Architecture split.")
    add_bullet(doc, ".NET Aspire for local multi-service run and test.")
    add_bullet(doc, "Blazor UI against those APIs.")
    add_bullet(doc, "Fewer unnecessary database round trips with a clearer command/query split.")

    add_para(doc, "sfCFO", 11, True, space_before=8)
    add_para(doc, "Jan 2026 – Feb 2026 · Remote", 10)
    add_bullet(doc, "Short scoped delivery on a .NET 9 Blazor Server product (MudBlazor, EF Core): Mapbox GL map view and E.164 phone input.")

    add_para(doc, "Policy Q&A", 11, True, space_before=8)
    add_para(doc, "2026 – Present · personal product", 10)
    add_bullet(doc, "Markdown policies to retrieval to answers with citations.")
    add_bullet(doc, "C# Allow / Deny / NeedsHuman layer, with tests.")
    add_bullet(doc, "Blazor, .NET Aspire, Ollama.")
    add_bullet(doc, "https://github.com/moserkristian/Moser.RagAi")

    add_para(doc, "Senior .NET engineer — Thermo Fisher Scientific", 11, True, space_before=14)
    add_para(doc, "Sep 2022 – Apr 2024 · Bratislava · Hybrid", 10)
    add_bullet(doc, "ASP.NET Core Web APIs for laboratory test results.")
    add_bullet(doc, "AWS for file handling.")
    add_bullet(doc, "Tests, monitoring, pipelines.")
    add_bullet(doc, "Written English across time zones.")

    add_para(doc, "Software engineer — Sociálna poisťovňa", 11, True, space_before=8)
    add_para(doc, "Sep 2019 – Jan 2022 · Bratislava · Hybrid", 10)
    add_bullet(doc, ".NET internal systems.")
    add_bullet(doc, "Incident handling.")
    add_bullet(doc, "Backup and recovery automation.")

    add_para(doc, "Junior software engineer — Frequentis", 11, True, space_before=8)
    add_para(doc, "Jan 2018 – Jun 2019 · Bratislava · Hybrid", 10)
    add_bullet(doc, "Optimised 6–7 multithreaded components; measured about 24% better performance.")
    add_bullet(doc, "Real-time client–server UI.")
    add_bullet(doc, "Small HTML/JS client.")

    add_para(doc, "Languages", 12, True, space_before=10, space_after=4)
    add_para(doc, "Slovak (native), English (professional).", 11)

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
